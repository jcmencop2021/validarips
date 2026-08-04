from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from models.relation_record import RELATION_COLUMNS, RelationRecord

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="000000")
HEADER_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(bold=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def _style_header_row(ws, row: int = 1) -> None:
    for col_idx, _ in enumerate(RELATION_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx)
        cell.fill = copy(HEADER_FILL)
        cell.border = copy(HEADER_BORDER)
        cell.font = copy(HEADER_FONT)
        cell.alignment = copy(HEADER_ALIGN)


def ensure_template(template_path: Path) -> Path:
    """Crea la plantilla Excel si no existe."""
    template_path.parent.mkdir(parents=True, exist_ok=True)
    if template_path.is_file():
        return template_path

    wb = Workbook()
    ws = wb.active
    ws.title = "Relacion"
    for col_idx, name in enumerate(RELATION_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    _style_header_row(ws)
    wb.save(template_path)
    return template_path


def export_to_excel(
    records: list[RelationRecord],
    output_path: Path,
    template_path: Path,
) -> None:
    ensure_template(template_path)
    wb = load_workbook(template_path)
    ws = wb.active

    # Limpiar filas de datos previas (mantener encabezado)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_idx, record in enumerate(records, start=2):
        for col_idx, value in enumerate(record.to_row(), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
