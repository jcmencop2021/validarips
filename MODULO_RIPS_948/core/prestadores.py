from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

NIT_ALIASES = frozenset(
    {
        "nit",
        "nits",
        "numdocumentoidobligado",
        "documento",
        "identificacion",
        "numero_identificacion",
        "codigo",
        "id",
        "numnit",
    }
)
NOMBRE_ALIASES = frozenset(
    {
        "nombre",
        "nombreips",
        "nombre_ips",
        "razonsocial",
        "razon_social",
        "prestador",
        "nombreprestador",
        "entidad",
        "descripcion",
        "nombreentidad",
        "nombrerazonsocial",
    }
)


def _norm_header(value: str) -> str:
    text = re.sub(r"[^a-z0-9]", "", str(value or "").lower())
    return text


def _norm_nit(value: str) -> str:
    return re.sub(r"\D", "", str(value or ""))


class PrestadoresCatalog:
    def __init__(self) -> None:
        self._by_nit: dict[str, str] = {}
        self.source_file: str = ""
        self.rows_loaded: int = 0

    def __len__(self) -> int:
        return len(self._by_nit)

    def get(self, nit: str) -> str:
        n = _norm_nit(nit)
        if not n:
            return ""
        candidates = [n]
        if len(n) > 9:
            candidates.append(n[:9])
        if len(n) > 1:
            candidates.append(n[:-1])
        for key in candidates:
            if key in self._by_nit:
                return self._by_nit[key]
        return ""

    def load_file(self, path: Path) -> int:
        self._by_nit.clear()
        self.rows_loaded = 0
        suffix = path.suffix.lower()
        if suffix in (".xlsx", ".xlsm"):
            count = self._load_xlsx(path)
        elif suffix == ".csv":
            count = self._load_csv(path, delimiter=",")
        elif suffix == ".txt":
            count = self._load_csv(path, delimiter=None)
        else:
            return 0
        self.rows_loaded = count
        if count:
            self.source_file = str(path)
        return count

    def _register_row(self, nit: str, nombre: str) -> bool:
        n = _norm_nit(nit)
        name = str(nombre or "").strip()
        if not n or not name:
            return False
        self._by_nit[n] = name
        return True

    def _detect_columns(self, headers: list[str]) -> tuple[int | None, int | None]:
        nit_idx: int | None = None
        name_idx: int | None = None
        for i, h in enumerate(headers):
            norm = _norm_header(h)
            if nit_idx is None and norm in NIT_ALIASES:
                nit_idx = i
            if name_idx is None and norm in NOMBRE_ALIASES:
                name_idx = i
        if nit_idx is None:
            for i, h in enumerate(headers):
                if "nit" in _norm_header(h) or "documento" in _norm_header(h):
                    nit_idx = i
                    break
        if name_idx is None:
            for i, h in enumerate(headers):
                norm = _norm_header(h)
                if any(x in norm for x in ("nombre", "razon", "prestador", "entidad")):
                    name_idx = i
                    break
        if nit_idx is None and name_idx is None and len(headers) >= 2:
            nit_idx, name_idx = 0, 1
        return nit_idx, name_idx

    def _load_xlsx(self, path: Path) -> int:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return 0
        headers = [str(c or "") for c in header_row]
        nit_idx, name_idx = self._detect_columns(headers)
        if nit_idx is None or name_idx is None:
            return 0
        count = 0
        for row in rows:
            if not row:
                continue
            nit = row[nit_idx] if nit_idx < len(row) else ""
            nombre = row[name_idx] if name_idx < len(row) else ""
            if self._register_row(str(nit), str(nombre)):
                count += 1
        return count

    def _load_csv(self, path: Path, delimiter: str | None) -> int:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        if delimiter is None:
            delimiter = ";" if text.count(";") > text.count(",") else ","
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        try:
            headers = next(reader)
        except StopIteration:
            return 0
        nit_idx, name_idx = self._detect_columns(headers)
        if nit_idx is None or name_idx is None:
            return 0
        count = 0
        for row in reader:
            if not row:
                continue
            nit = row[nit_idx] if nit_idx < len(row) else ""
            nombre = row[name_idx] if name_idx < len(row) else ""
            if self._register_row(nit, nombre):
                count += 1
        return count


def find_prestadores_file(search_roots: set[Path]) -> Path | None:
    extensions = {".xlsx", ".xlsm", ".csv", ".txt"}
    candidates: list[Path] = []

    def scan_docs(base: Path) -> None:
        docs = base / "docs"
        if not docs.is_dir():
            return
        for path in sorted(docs.iterdir()):
            if not path.is_file() or path.suffix.lower() not in extensions:
                continue
            name = path.name.lower()
            if "prestador" in name or name.startswith("entidad"):
                candidates.insert(0, path)
            else:
                candidates.append(path)

    for root in search_roots:
        if root.is_dir():
            scan_docs(root)
            scan_docs(root.parent)

    if not candidates:
        return None

    def sort_key(p: Path) -> tuple[int, str]:
        n = p.name.lower()
        priority = 0 if "prestador" in n else 1
        return (priority, n)

    return sorted(candidates, key=sort_key)[0]


def load_prestadores_catalog(search_roots: set[Path]) -> PrestadoresCatalog:
    catalog = PrestadoresCatalog()
    path = find_prestadores_file(search_roots)
    if path:
        catalog.load_file(path)
    return catalog
